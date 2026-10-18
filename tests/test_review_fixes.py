import pytest
from openpyxl import load_workbook

from jobpilot import server
from jobpilot.config import StructuredConfig
from jobpilot.joblist import write_job_list
from jobpilot.sources.jsonld_source import JsonLdSource
from jobpilot.store.db import Store
from jobpilot.store.models import Posting

PROPERTY_VALUE_HTML = """
<script type="application/ld+json">
{"@context":"https://schema.org","@graph":[
 {"@type":"JobPosting","identifier":{"@type":"PropertyValue","name":"Acme Corp","value":"J-100"},
  "title":"Analyst","url":"https://careers.example.test/1","hiringOrganization":{"name":"Acme Corp"},
  "jobLocation":{"address":{"addressLocality":"Austin"}},"description":"d"},
 {"@type":"JobPosting","identifier":{"@type":"PropertyValue","name":"Acme Corp","value":"J-200"},
  "title":"Senior Analyst","url":"https://careers.example.test/2","hiringOrganization":{"name":"Acme Corp"},
  "jobLocation":{"address":{"addressLocality":"Austin"}},"description":"d"}]}
</script>
"""


def test_jsonld_property_value_ids_do_not_collapse_to_the_company_name():
    cfg = StructuredConfig(name="acme", company_name="Acme Corp", endpoint="https://careers.example.test",
                           enabled=True, canary_min_results=1)
    found, raw = JsonLdSource(cfg, fetcher=lambda url, timeout: PROPERTY_VALUE_HTML).search("analyst", None)
    assert raw == 2
    assert sorted(p.external_id for p in found) == ["J-100", "J-200"]


@pytest.mark.parametrize("evil", ["=HYPERLINK(\"http://evil.test\",\"x\")", "+SUM(1)", "-2+3", "@cmd"])
def test_job_list_never_contains_live_formulas(tmp_path, evil):
    store = Store(tmp_path / "j.db")
    store.migrate()
    profile = tmp_path / "p" / "profile.yaml"
    profile.parent.mkdir()
    profile.write_text("name: A\nheadline: M\nlocation: X\nsources: [indeed_in]\nsearch_terms: [m]\n", encoding="utf-8")
    sid = store.add_seeker("A", None, str(profile), 20)
    pid = store.upsert_posting(Posting(company_name=evil, source="indeed_in", external_id="1", title=evil,
                                       location=evil, description="d", url="https://jobs.test/1", posted_at="2026-10-18"))
    store.link_posting(sid, pid)
    store.record_score(sid, pid, 90, evil, "t")
    write_job_list(store, sid, tmp_path / "out.xlsx", min_score=60)
    ws = load_workbook(tmp_path / "out.xlsx").active
    for col in (2, 3, 4, 8):
        cell = ws.cell(2, col)
        assert cell.data_type != "f" and str(cell.value).startswith("'"), (col, cell.value)
    store.close()


def test_empty_onboarding_variable_refuses_to_open_the_operator_tools(monkeypatch):
    monkeypatch.delenv("JOBPILOT_SEEKER_ID", raising=False)
    monkeypatch.setenv("JOBPILOT_ONBOARD_TELEGRAM_ID", "")
    with pytest.raises(SystemExit, match="empty"):
        server.main()


def test_profile_writes_are_atomic_and_leave_no_temp_files(tmp_path):
    f = tmp_path / "profile.yaml"
    f.write_text("old", encoding="utf-8")
    server._atomic_write(f, "new: 1\n")
    assert f.read_text() == "new: 1\n" and not list(tmp_path.glob("*.tmp"))


def test_one_broken_profile_does_not_stop_discovery_for_others(tmp_path, monkeypatch):
    monkeypatch.setenv("JOBPILOT_HOME", str(tmp_path))
    (tmp_path / "config.toml").write_text(
        '[paths]\ndata_dir="data"\nout_dir="out"\n[limits]\nscore_batch=40\nmin_request_gap_seconds=4.0\n', encoding="utf-8")
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "companies.csv").write_text("name,tier,ats_vendor,careers_url,region,enabled\n", encoding="utf-8")
    store = Store(tmp_path / "j.db")
    store.migrate()
    good = tmp_path / "data" / "seekers" / "g"
    good.mkdir(parents=True)
    (good / "profile.yaml").write_text("name: G\nheadline: PM\nlocation: X\nsources: [none]\nsearch_terms: [pm]\n", encoding="utf-8")
    store.add_seeker("Good", 1, "data/seekers/g/profile.yaml", 20)
    store.add_seeker("Broken", 2, "data/seekers/missing/profile.yaml", 20)
    tools = server.build_tools(store, tmp_path / "out", operator_mode=True)
    out = tools["discover"]()
    assert out["ok"] is True and 2 in out["skipped_seekers"] and 1 not in out["skipped_seekers"]
    store.close()
