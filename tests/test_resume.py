
import pytest

from jobpilot.resume import (
    ResumeValidationError,
    load_base_resume,
    render_docx,
    validate_docx,
    validate_pdf,
)


BASE = """
name: Ada Example
headline: Senior Product Merchandiser
contact: [ada@example.test, +1 555 0100]
summary: Product and sourcing professional with Ralph Lauren account experience.
required_facts:
  - Ada Example
  - Senior Product Merchandiser
  - Example Brands Ltd
  - Northwind Buying Office
  - Ralph Lauren
experience:
  - id: current
    employer: Example Brands Ltd
    title: Senior Product Merchandiser
    dates: 2022 - Present
    bullets:
      - id: current-1
        text: Owned product development from concept to delivery.
      - id: current-2
        text: Managed supplier capability and critical path execution.
  - id: prior
    employer: Northwind Buying Office
    title: Merchandiser
    dates: 2019 - 2022
    bullets:
      - id: prior-1
        text: Coordinated vendor communication and sample approvals.
education:
  - id: degree
    institution: Example University
    degree: Bachelor of Design
    dates: 2015 - 2019
skills:
  - Product development
  - Supplier management
  - Critical path ownership
"""


@pytest.fixture
def base_resume(tmp_path):
    path = tmp_path / "base_resume.yaml"
    path.write_text(BASE, encoding="utf-8")
    return load_base_resume(path)


def test_load_base_resume_requires_structured_facts(base_resume):
    assert base_resume.name == "Ada Example"
    assert [item.id for item in base_resume.experience] == ["current", "prior"]
    assert base_resume.experience[0].bullets[0][0] == "current-1"
    assert "Ralph Lauren" in base_resume.required_facts


def test_render_docx_is_single_column_and_passes_extraction_gate(base_resume, tmp_path):
    output = tmp_path / "resume.docx"

    report = render_docx(base_resume, output)

    assert report.ok is True
    assert report.format == "docx"
    assert output.exists()
    extracted = validate_docx(output, base_resume.required_facts)
    assert extracted.ok is True


def test_render_can_reorder_only_known_facts(base_resume, tmp_path):
    report = render_docx(
        base_resume,
        tmp_path / "resume.docx",
        experience_order=["prior", "current"],
        skills=["Supplier management", "Product development"],
    )

    assert report.ok


def test_unknown_selection_is_rejected_instead_of_invented(base_resume, tmp_path):
    with pytest.raises(ValueError, match="unknown facts"):
        render_docx(base_resume, tmp_path / "resume.docx", experience_order=["invented"])

    with pytest.raises(ValueError, match="not present"):
        render_docx(base_resume, tmp_path / "resume.docx", skills=["Invented skill"])


def test_missing_required_fact_blocks_delivery(base_resume, tmp_path):
    base = type(base_resume)(
        name=base_resume.name,
        headline=base_resume.headline,
        contact=base_resume.contact,
        summary=base_resume.summary,
        experience=base_resume.experience,
        education=base_resume.education,
        skills=base_resume.skills,
        required_facts=(*base_resume.required_facts, "Fact absent from source"),
    )
    output = tmp_path / "blocked.docx"

    with pytest.raises(ResumeValidationError) as exc_info:
        render_docx(base, output)

    assert "Fact absent from source" in exc_info.value.report.missing
    assert not output.exists()


def test_plain_string_bullets_get_stable_ids(tmp_path):
    path = tmp_path / "base.yaml"
    path.write_text(
        "name: A\nheadline: H\nexperience:\n"
        "  - id: e\n    employer: E\n    title: T\n    dates: D\n"
        "    bullets: [One factual bullet]\n",
        encoding="utf-8",
    )

    base = load_base_resume(path)
    assert base.experience[0].bullets == (("e-bullet-1", "One factual bullet"),)


def test_pdf_extraction_gate_reads_text(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    output = tmp_path / "resume.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=200)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    resources = DictionaryObject({NameObject("/Font"): DictionaryObject({
        NameObject("/F1"): writer._add_object(font),
    })})
    page[NameObject("/Resources")] = resources
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 20 150 Td (Ada Example) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    with output.open("wb") as handle:
        writer.write(handle)

    report = validate_pdf(output, ("Ada Example",))
    assert report.ok is True


def test_unverified_facts_tolerates_pdf_spacing_but_not_new_claims(base_resume):
    from jobpilot.resume import unverified_facts
    source = BASE.replace("Owned product development", "Owned   product\ndevelop ment")
    assert unverified_facts(base_resume, source) == []
    assert unverified_facts(base_resume, source.replace("Example Brands Ltd", "")) == [
        "Example Brands Ltd"]


def test_extract_source_text_reads_docx_tables(tmp_path):
    from docx import Document
    from jobpilot.resume import extract_source_text
    doc = Document()
    doc.add_paragraph("Ada Example")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Example Brands Ltd"
    doc.save(tmp_path / "cv.docx")
    text = extract_source_text(tmp_path / "cv.docx")
    assert "Ada Example" in text and "Example Brands Ltd" in text
    with pytest.raises(ValueError):
        extract_source_text(tmp_path / "cv.png")
