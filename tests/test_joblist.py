from openpyxl import load_workbook

from jobpilot import telegram_bot as tb
from jobpilot.joblist import write_job_list
from jobpilot.notify import push_job_lists
from jobpilot.store.db import Store
from jobpilot.store.models import Posting


def _post(store, sid, n, source="indeed_in", title=None, url="https://jobs.test/{n}"):
    pid = store.upsert_posting(Posting(
        company_name=f"Co{n}", source=source, external_id=str(n), title=title or f"Role {n}",
        location="Gurgaon", description="d", url=url.format(n=n), posted_at="2026-09-20T00:00:00"))
    store.link_posting(sid, pid)
    return pid


def _seeker(tmp_path, store, uid=None):
    profile = tmp_path / "p" / "profile.yaml"
    profile.parent.mkdir(exist_ok=True)
    profile.write_text("name: A\nheadline: M\nlocation: Gurgaon\nsources: [indeed_in]\n"
                       "search_terms: [m]\nmin_score: 60\n", encoding="utf-8")
    return store.add_seeker("A", uid, str(profile), 20)


def test_list_is_scored_best_first_then_unrated_with_links(tmp_path):
    store = Store(tmp_path / "j.db"); store.migrate()
    sid = _seeker(tmp_path, store)
    low, high, _unrated = (_post(store, sid, n) for n in (1, 2, 3))
    store.record_score(sid, low, 40, "weak fit")
    store.record_score(sid, high, 90, "strong fit")
    total, good = write_job_list(store, sid, tmp_path / "out.xlsx", min_score=60)
    rows = list(load_workbook(tmp_path / "out.xlsx").active.iter_rows(values_only=True))
    assert (total, good) == (3, 1)
    assert rows[0][0] == "Fit score"
    assert [r[0] for r in rows[1:]] == [90, 40, None]
    assert rows[3][7] == "Not rated yet"
    cell = load_workbook(tmp_path / "out.xlsx").active.cell(2, 10)
    assert cell.value == "Open" and cell.hyperlink.target == "https://jobs.test/2"


def test_same_role_on_two_boards_is_one_row(tmp_path):
    store = Store(tmp_path / "j.db"); store.migrate()
    sid = _seeker(tmp_path, store)
    _post(store, sid, 1, source="indeed_in", title="Merchandiser")
    store.link_posting(sid, store.upsert_posting(Posting(
        company_name="Co1", source="foundit_in", external_id="x", title="Merchandiser",
        location="Gurgaon", description="d", url="https://f.test/x", posted_at=None)))
    assert write_job_list(store, sid, tmp_path / "o.xlsx")[0] == 1
    found_on = load_workbook(tmp_path / "o.xlsx").active.cell(2, 7).value
    assert set(found_on.split(", ")) == {"indeed_in", "foundit_in"}


def _bot(tmp_path, store):
    sent, files = [], []
    bot = tb.Bot("T", tmp_path, store, http=lambda m, p, timeout=40: sent.append((m, p)) or {"ok": True},
                 runner=None, upload=lambda chat, path: files.append(path))
    return bot, sent, files


def test_list_command_sends_the_file_without_claude(tmp_path):
    store = Store(tmp_path / "j.db"); store.migrate()
    sid = _seeker(tmp_path, store, uid=5)
    _post(store, sid, 1)
    bot, sent, files = _bot(tmp_path, store)
    bot.handle({"update_id": 1, "message": {"from": {"id": 5, "is_bot": False},
                                            "chat": {"id": 5, "type": "private"}, "text": "/list"}})
    assert len(files) == 1 and files[0].suffix == ".xlsx"
    assert "1 jobs" in [p["text"] for m, p in sent if m == "sendMessage"][-1]


def test_daily_list_goes_only_to_seekers_with_new_jobs(tmp_path):
    store = Store(tmp_path / "j.db"); store.migrate()
    sid = _seeker(tmp_path, store, uid=5)
    bot, _, files = _bot(tmp_path, store)
    assert push_job_lists(store, bot.send_job_list) == 0 and files == []
    _post(store, sid, 1)
    assert push_job_lists(store, bot.send_job_list) == 1 and len(files) == 1
