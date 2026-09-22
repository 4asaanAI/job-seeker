"""One consolidated spreadsheet of every job found for a seeker.

Each row is one real-world role (boards carrying the same role are already
merged), best fit first; jobs not yet rated follow, newest first.
"""
from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path

from .sources.base import safe_url
from .store.db import Store

MAX_AGE_DAYS = 30
COLUMNS = (("Fit score", 10), ("Job title", 42), ("Company", 26), ("Location", 26),
           ("Posted", 12), ("First seen", 12), ("Found on", 20), ("Why it fits", 70),
           ("My status", 14), ("Link", 10))


def job_rows(store: Store, seeker_id: int, days: int = MAX_AGE_DAYS) -> list[dict]:
    """Jobs still seen within `days`, scored first (high to low), then unrated."""
    since = (datetime.now(timezone.utc) - timedelta(days=days)).isoformat()
    with store._lock:
        rows = store.conn.execute(
            """
            SELECT p.id, p.title, p.company_name, p.location, p.posted_at,
                   sj.first_seen_at, p.canonical_url, s.fit_score, s.reasoning,
                   a.status AS applied,
                   (SELECT group_concat(DISTINCT ps.source) FROM posting_sources ps
                     WHERE ps.posting_id = p.id) AS sources,
                   (SELECT ps.url FROM posting_sources ps
                     WHERE ps.posting_id = p.id AND ps.url IS NOT NULL LIMIT 1) AS source_url
            FROM seeker_jobs sj
            JOIN job_postings p ON p.id = sj.posting_id
            LEFT JOIN scores s ON s.seeker_id = sj.seeker_id AND s.posting_id = p.id
            LEFT JOIN applications a ON a.seeker_id = sj.seeker_id AND a.posting_id = p.id
            WHERE sj.seeker_id = ? AND p.last_seen_at >= ?
            ORDER BY s.fit_score IS NULL, s.fit_score DESC, sj.first_seen_at DESC
            """,
            (seeker_id, since),
        ).fetchall()
    return [dict(r) for r in rows]


def write_job_list(store: Store, seeker_id: int, path: Path,
                   min_score: float = 70.0) -> tuple[int, int]:
    """Write the spreadsheet; returns (jobs, good matches)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    rows = job_rows(store, seeker_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    ws.append([name for name, _ in COLUMNS])
    for i, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
        ws.cell(1, i).font = Font(bold=True, color="FFFFFF")
        ws.cell(1, i).fill = PatternFill("solid", fgColor="305496")
    good_fill = PatternFill("solid", fgColor="C6EFCE")
    good = 0
    for r in rows:
        score = r["fit_score"]
        url = safe_url(r["canonical_url"] or r["source_url"] or "")
        ws.append([
            round(score) if score is not None else None,
            r["title"], r["company_name"], r["location"] or "",
            (r["posted_at"] or "")[:10], (r["first_seen_at"] or "")[:10],
            (r["sources"] or "").replace(",", ", "),
            r["reasoning"] if score is not None else "Not rated yet",
            r["applied"] or "", "Open" if url else "",
        ])
        n = ws.max_row
        if url:
            link = ws.cell(n, len(COLUMNS))
            link.hyperlink = url
            link.font = Font(color="0563C1", underline="single")
        ws.cell(n, 8).alignment = Alignment(wrap_text=True, vertical="top")
        if score is not None and score >= min_score:
            good += 1
            ws.cell(n, 1).fill = good_fill
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return len(rows), good
